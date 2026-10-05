"""Excel reporting for batch visibility results."""

from __future__ import annotations

from collections import Counter
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = "1F4E78"
HEADER_FONT = "FFFFFF"
SUBTLE_FILL = "D9EAF7"


def _format_sheet(ws) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.font = Font(color=HEADER_FONT, bold=True)
        cell.alignment = Alignment(horizontal="center")
    for column in ws.columns:
        width = min(max(len(str(c.value or "")) for c in column) + 2, 55)
        ws.column_dimensions[get_column_letter(column[0].column)].width = width


def _write_table(ws, rows: Iterable[dict], columns: list[str]) -> None:
    ws.append(columns)
    for row in rows:
        ws.append([row.get(col) for col in columns])
    _format_sheet(ws)


def write_report(
    output_xlsx: str | Path,
    site_summary: list[dict],
    visibility_rows: list[dict],
    qa_rows: list[dict],
    config: dict,
) -> Path:
    """Create the final Excel workbook."""
    output_xlsx = Path(output_xlsx)
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"

    statuses = Counter(row.get("Status", "Unknown") for row in site_summary)
    visible_count = sum(int(row.get("Visible_Towers", 0) or 0) for row in site_summary)
    observer_count = sum(int(row.get("Observer_Count", 0) or 0) for row in site_summary)

    dashboard = [
        ("Viewshed & Visibility Automation", ""),
        ("Sites submitted", len(site_summary)),
        ("Sites complete", statuses.get("Complete", 0)),
        ("Sites failed", statuses.get("Failed", 0)),
        ("Observers generated", observer_count),
        ("Visible site-to-tower relationships", visible_count),
    ]
    for label, value in dashboard:
        ws.append([label, value])
    ws["A1"].font = Font(size=16, bold=True)
    ws["A1"].fill = PatternFill("solid", fgColor=HEADER_FILL)
    ws["A1"].font = Font(color=HEADER_FONT, size=16, bold=True)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 22
    for row in range(2, len(dashboard) + 1):
        ws[f"A{row}"].fill = PatternFill("solid", fgColor=SUBTLE_FILL)
        ws[f"A{row}"].font = Font(bold=True)

    summary_ws = wb.create_sheet("Site Summary")
    _write_table(
        summary_ws,
        site_summary,
        [
            "Site_ID",
            "Geometry_Type",
            "Observer_Count",
            "Candidate_Towers",
            "Visible_Towers",
            "Obscured_Towers",
            "NoData_Towers",
            "Status",
            "Elapsed_Minutes",
            "Message",
        ],
    )

    detail_ws = wb.create_sheet("Visibility Detail")
    _write_table(
        detail_ws,
        visibility_rows,
        [
            "Site_ID",
            "Tower_ID",
            "Tower_Height_Input",
            "Tower_Height_Units",
            "Tower_Height_m",
            "Required_AGL_m",
            "Visibility",
        ],
    )

    qa_ws = wb.create_sheet("QA Report")
    _write_table(qa_ws, qa_rows, ["Level", "Check", "Result", "Message"])

    cfg_ws = wb.create_sheet("Configuration")
    cfg_ws.append(["Setting", "Value"])
    for key, value in sorted(config.items()):
        cfg_ws.append([key, str(value)])
    _format_sheet(cfg_ws)

    wb.save(output_xlsx)
    return output_xlsx
